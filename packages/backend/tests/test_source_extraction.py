from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.source_extraction import extract_documents


@pytest.mark.asyncio
async def test_pdf_preserves_empty_page_numbers(monkeypatch, tmp_path: Path):
    import pdfplumber

    class PDF:
        pages = [SimpleNamespace(extract_text=lambda: None),
                 SimpleNamespace(extract_text=lambda: "second page")]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

    monkeypatch.setattr(pdfplumber, "open", lambda _: PDF())
    result = await extract_documents(tmp_path / "sample.pdf")
    assert [item.logical_path for item in result] == ["page/1", "page/2"]
    assert result[0].text == ""
    assert result[1].locator["page"] == 2


@pytest.mark.asyncio
async def test_docx_keeps_empty_paragraph_and_whitespace(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("first")
    doc.add_paragraph("")
    doc.add_paragraph("  third  ")
    doc.save(str(path))
    result = await extract_documents(path)
    assert [item.text for item in result] == ["first", "", "  third  "]
    assert result[2].logical_path == "paragraph/3"
    assert result[2].locator["paragraph"] == 3


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", [".md", ".txt", ".py"])
async def test_plain_text_preserves_newlines_and_indent(tmp_path: Path, suffix: str):
    path = tmp_path / f"source{suffix}"
    raw = b"def run():\r\n    return 1\r\n"
    path.write_bytes(raw)
    result = await extract_documents(path)
    assert result[0].text == raw.decode()
    assert result[0].locator["char_end"] == len(raw)
    assert result[0].locator["line_start"] == 1
